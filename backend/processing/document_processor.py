import re
from pathlib import Path


def clean_text(text: str) -> str:
    text = re.sub(r"\r", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(path: Path) -> str:
    import fitz

    parts = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    text = clean_text("\n".join(parts))
    if not text:
        raise ValueError("The uploaded PDF did not contain extractable text")
    return text


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    text = clean_text("\n".join(parts))
    if not text:
        raise ValueError("The uploaded DOCX did not contain extractable text")
    return text


def extract_text_from_file(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))

    if suffix == ".pdf":
        return _extract_pdf(path)

    if suffix in {".docx", ".doc"}:
        return _extract_docx(path)

    raise ValueError(f"Unsupported file type '{suffix}'. Upload PDF, DOCX, or TXT.")
